"""Build the data access note as PDF and Word from one content source.

The note documents sources.py, network_tools.py, site_rules.py, crawler.py, web_links.py and
source_material.py. Keep the numbers here in step with those files and with config.py: every
limit quoted below is a setting or constant there. Its layout follows the role-play moderation
note (docs/roleplay-moderation/build.py), so the lab's notes read as one series.

The generators are not dependencies of the app, so run this from a throwaway environment:

    uv venv /tmp/docenv && uv pip install --python /tmp/docenv/bin/python python-docx reportlab
    /tmp/docenv/bin/python docs/network-access/build.py

Output goes to src/ailawlab/web/static/docs/, where the web portal links to it.
"""
from __future__ import annotations

import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "src" / "ailawlab" / "web" / "static" / "docs"
STEM = "network-access"

TITLE = "Outside Data Access and Agent Web Reading in the AI Law Lab"
SUBTITLE = "Which services the lab calls, when its agents go online, and how a site's wishes are honoured"
AUTHOR = "Jian Gong"
AFFILIATION = "AI Law Lab, University of Wyoming"
DATELINE = "Data access note, October 2026. Describes the software at commit fa0be1c."
FOOTER = "AI Law Lab · Data access note · Outside data access and agent web reading"


def t(text: str) -> str:
    """Collapse a triple-quoted paragraph to one line."""
    return " ".join(text.split())


ABSTRACT = t("""
The AI Law Lab runs its language models on university hardware and answers questions from its
own versioned libraries, but a research run may also search public legal databases and the
open web, and read what it finds. This note sets out every service the lab's server connects
to, what is sent to each, and the exact points at which a run goes online. It describes the
five legal-database interfaces the lab uses and how agents' calls to them are paced; how every
document read is saved into a library before it can be cited; and the checks a web page must
pass before a research agent or a crawl may read it. A page is read only if its site permits
it: robots.txt is obeyed for the lab's own name and for well-known AI agents, and pages marked
"noai" or carrying a text-and-data-mining reservation are left alone. Refusals are reported to
the agent and listed on the run's page. The note closes with the safeguards applied to every
fetch and the questions that remain open.
""")

# Blocks: ("h1", text) ("h2", text) ("p", text) ("table", spec) ("algo", spec) ("bullets", [text])
# ("refs", [text]). Inline markup: **bold**, *italic*, `code`. URLs are linked in the PDF.
CONTENT: list[tuple] = [
    ("h1", "1  Introduction"),
    ("p", t("""
A study of how language models reason about law should be able to say exactly what the model
was shown. The lab's runs therefore draw on libraries whose every change is recorded as a
numbered version, and each citation in an answer is traced to a document, a library and a
version. Reaching outside the libraries threatens both properties: a web page can change or
vanish after it is read, and a search sends part of the researcher's question to a third party.
""")),
    ("p", t("""
Outside access is nonetheless useful. Agency guidance, recent opinions and scholarship are
often not yet in any library. The lab allows it under three conditions. A person must allow
it, for each run. Whatever is read is saved into a library first, so that a citation to it
names a fixed copy. And a page is read only if its website has not said that it does not want
AI systems reading it. Section 2 lists every outside connection; Section 3 describes the legal
databases; Section 4 the points at which a research run goes online; Section 5 what people add
by hand; Section 6 the checks on a site's wishes; Section 7 the safeguards on every fetch; and
Section 9 the limitations.
""")),

    ("h1", "2  Every outside connection"),
    ("p", t("""
Table 1 lists every destination the lab's server connects to. Model inference never leaves
university hardware, and no prompt, document or answer is sent to a commercial model API. What
leaves the server is a search term or an address, never a document a researcher uploaded.
""")),
    ("table", {
        "caption": t("""Table 1. Every destination the server connects to, what is sent, and when.
            The search terms an agent writes are drawn from the research brief, so a brief that
            must stay confidential should be run with both outside permissions off."""),
        "header": ["Destination", "What is sent", "When"],
        "rows": [
            ["The two DGX Spark units, inside the university",
             "Prompts, and text to embed", "Every model call"],
            ["CourtListener, Federal Register, eCFR, govinfo, SEC EDGAR",
             "A search term, or the identifier of a result",
             "A person searches or adds on the Legal Sources page; an agent searches or reads "
             "when legal databases are allowed"],
            ["Brave Search API", "A search query written by an agent",
             "An agent searches the open web, when the open web is allowed"],
            ["Any public website", "A request for one address",
             "An agent reads a page; a page named in the brief is read; a crawl; a person adds "
             "or re-checks a link, or drafts a cast from one"],
            ["UW Shibboleth, InCommon metadata service, Google, Microsoft",
             "Sign-in exchanges, through Keycloak", "Someone signs in with that provider"],
        ],
        "widths": [0.30, 0.27, 0.43],
    }),

    ("h1", "3  The legal databases"),
    ("h2", "3.1  Interfaces and credentials"),
    ("p", t("""
Five public repositories are reached through the interfaces their publishers provide for
programs to use (Table 2). Each is a staging step rather than a live oracle: a search returns
candidates, and nothing is relied on until it has been fetched, divided into passages, embedded
and saved into a library with its canonical address and the time it was retrieved. Credentials
are kept in the server's environment, never in the code or the browser, and a database whose
credential is missing reports itself unavailable instead of failing.
""")),
    ("table", {
        "caption": "Table 2. The legal databases, the interface used for each, and what it returns.",
        "header": ["Database", "Interface and credential", "What it holds"],
        "rows": [
            ["CourtListener", "REST API v4 [5]; API token, set on the server",
             "Opinions of federal and state courts (Free Law Project). Without a token only the "
             "search snippet can be saved, and it is labelled as a snippet."],
            ["Federal Register", "API v1 [6]; no key",
             "Proposed and final rules, notices and presidential documents since 1994; full text "
             "where published, else the abstract, flagged as such."],
            ["eCFR", "Search and versioner APIs [7]; no key",
             "The Code of Federal Regulations as in force, read as of each title's latest issue."],
            ["govinfo (GPO)", "govinfo API [8]; api.data.gov key, set on the server",
             "US Code, Public Laws, bills and reports; HTML or text rendition, else the PDF."],
            ["SEC EDGAR", "Full-text search and the EDGAR archives; no key",
             "Company filings since 2001. Requests name the lab and a contact address, as the "
             "SEC's fair-access policy requires [9]."],
        ],
        "widths": [0.17, 0.33, 0.50],
    }),
    ("h2", "3.2  Pacing"),
    ("p", t("""
A person's search on the Legal Sources page goes straight through. Research agents call far
more often, several at a time, so each database's requests from agents are spaced at least one
second apart across every run on the server. A database that answers HTTP 429 ("too many
requests") is left alone for three minutes, and agents are told to search elsewhere meanwhile.
One agent search asks each database for four results.
""")),
    ("h2", "3.3  Open-web search"),
    ("p", t("""
Brave's Search API [10] answers an agent's web searches, eight results each. Calls are spaced
1.1 seconds apart, the API's per-second limit, and one 429 is retried after two seconds.
Results are numbered leads ([W1], [W2], ...): an agent cannot cite a search result, only what
it has read.
""")),
    ("h2", "3.4  Why robots.txt does not govern these"),
    ("p", t("""
robots.txt addresses crawlers that wander a website. These interfaces exist so that programs can
query them; their use is governed by their own terms, keys and rate limits, which the lab
follows. The documents they return are largely government works [11] or public filings. That a
document is public does not settle every use of it, and researchers remain responsible for
theirs.
""")),

    ("h1", "4  When a research run goes online"),
    ("h2", "4.1  Two permissions"),
    ("p", t("""
An agentic-workflow run has two permissions, each the researcher's to grant on the plan before
the run starts and both off by default: *search the legal databases*, and *search the open web*,
which includes reading any public page. With both off the run is entirely local. Role-play and
document-analysis runs never go online.
""")),
    ("h2", "4.2  The four stages of a run"),
    ("p", t("""
With either permission on, network access happens at these points and no others:
""")),
    ("bullets", [
        t("""**Planning: none.** The planner drafts sub-questions from the brief on the Spark units.
           It notes web addresses written in the brief but does not open them."""),
        t("""**Before research: pages named in the brief.** If the open web is allowed, each address
           in the brief is read once, before any agent starts, so that every agent has it. These
           reads pass the same checks as an agent's own (Section 6)."""),
        t("""**During research: the agents' tools.** Each research agent, up to four at once,
           decides for itself when to call `search_databases`, `search_web` and `read`. Reading is
           not budgeted; the run's time limit (two hours by default) and the researcher's Stop
           button end it. A document is read once per run, however many agents ask for it."""),
        t("""**Write-up: none.** The lead agent writes the answer from the findings and the passages
           already read."""),
    ]),
    ("h2", "4.3  How a read becomes a source"),
    ("p", t("""
Everything read is saved into a library, by default "Fetched: <experiment name>", which records
a new version; the run is then pinned to that version, so later library searches by any agent
cover the new document. The passages handed back to the agent are numbered in the run's ledger
like any other, so a citation to something read online names a document, a library and a
version, and stays readable after the page changes. The run's page lists every search with the
results it gave, every document read with the library version it was saved to, and every page
that was not read because its site does not allow it.
""")),

    ("h1", "5  What a person adds"),
    ("p", t("""
On the Legal Sources page a person can search the databases and add chosen results; add web
pages by address, up to 20 at a time, and later check them again, which re-reads each page and
rebuilds its passages only if the text has changed; or crawl a page's links into a library. In
the cast builder a person can draft a role-play from a page given by its address. Adding a link
or drafting from one is that person reading the page, as they would in a browser, so it is not
put through the checks of Section 6. A crawl, which reads many pages automatically, is.
""")),

    ("h1", "6  Honouring a site's wishes"),
    ("p", t("""
Before a research agent reads a page, and before each page of a crawl, the lab asks what the
site has published about automated and AI use, and obeys it. Procedure 1 gives the order for
an agent; a crawl asks the same questions under its own name.
""")),
    ("algo", {
        "caption": t("""Procedure 1. How a research agent's request to read a page is decided. Each
            refusal is returned to the agent with its reason, recorded in the trace and listed on
            the run's page."""),
        "lines": [
            "read(address):",
            "  if the address is not public http(s)            -> refuse  (unsafe; nothing fetched)",
            "  rules <- robots.txt and tdmrep.json of the site    (cached up to one hour)",
            "  if robots.txt cannot be read (site failing)    -> refuse",
            "  if robots.txt disallows AILawLab-agent here    -> refuse",
            "  if robots.txt disallows a named AI agent here  -> refuse",
            "  if tdmrep.json reserves TDM rights here         -> refuse",
            "  page <- fetch once as AILawLab-agent/1.0        (at most 8 MB, 45 s)",
            "  if the page moved to another site, check that site's rules too",
            "  if page headers or <meta> say noai,",
            "     or tdm-reservation: 1                        -> refuse  (fetched, not kept)",
            "  save page to the run's library; return its passages, numbered for citation",
        ],
    }),
    ("h2", "6.1  robots.txt, for the lab"),
    ("p", t("""
The site's robots.txt is read first, following the Robots Exclusion Protocol [1]. Research
agents are known to it as `AILawLab-agent` and crawls as `AILawLab-crawler`; a site may write
rules for either, or for all crawlers (`*`). The most specific matching rule decides, and Allow
wins a tie. The `*` and `$` wildcards defined by the protocol are understood; Python's built-in
robots.txt reader ignores them, which would have let a rule such as `Disallow: /*.pdf$` through.
A missing robots.txt sets no rules; one that cannot be read because the site is failing means
nothing is read, as the protocol prescribes.
""")),
    ("h2", "6.2  robots.txt, for AI agents"),
    ("p", t("""
Many sites now turn named AI agents away while admitting search engines. If a site's robots.txt
closes the page to any of the agents in Table 3, the lab treats that as a refusal of AI systems
in general, which it is one of. Only rules written for those agents by name count: an agent
that falls back to the rules for all crawlers is treated as the lab already is.
""")),
    ("table", {
        "caption": t("""Table 3. The AI agents whose exclusion from a page the lab honours as its own
            (`site_rules.AI_AGENTS`). The list is maintained by hand."""),
        "header": ["Kind", "Agents"],
        "rows": [
            ["Assistants reading for a person",
             "ChatGPT-User, Claude-User, Perplexity-User, MistralAI-User"],
            ["Crawlers of AI companies",
             "GPTBot, OAI-SearchBot, ClaudeBot, Claude-SearchBot, anthropic-ai, PerplexityBot, "
             "CCBot, Google-Extended, Applebot-Extended, Bytespider, meta-externalagent, cohere-ai"],
        ],
        "widths": [0.28, 0.72],
    }),
    ("h2", "6.3  Text-and-data-mining reservations"),
    ("p", t("""
The TDM Reservation Protocol [2] lets a rights holder reserve text-and-data-mining rights in a
machine-readable way, the kind of opt-out European copyright law recognizes [3]. A site may
publish the reservation in `/.well-known/tdmrep.json`, by location; a page may carry it in a
`tdm-reservation: 1` response header or a `tdm-reservation` meta tag. A reservation means the
page is not read, and the licensing address the site gives, if any, is passed on to the agent.
""")),
    ("h2", "6.4  \"noai\" on the page"),
    ("p", t("""
Some sites mark pages "noai", in an `X-Robots-Tag` response header or a robots meta tag in the
page's head. The lab honours the directive when it is addressed to every agent or to the lab by
name. These markings can be seen only once the page has arrived, so such a page is fetched but
neither kept nor shown to the agent.
""")),
    ("h2", "6.5  When a page is refused"),
    ("p", t("""
A refused agent is told plainly why and asked to find another source; the refusal is recorded
in the run's trace and listed, with its reason, on the run's page. A crawl records the page as
"not open to AI use" and carries on with the rest. A site's rules are read at most once an hour
per site, across all runs (the protocol allows a day), and a failing site is asked again next
time rather than remembered.
""")),

    ("h1", "7  Safeguards on every fetch"),
    ("bullets", [
        t("""**Public addresses only.** The server can reach the Spark units, its database and the
           university network, so a link is fetched only if it uses `http` or `https` and its host
           resolves exclusively to public addresses. Every redirect, up to five, is checked
           again."""),
        t("""**Bounded.** A response is cut off at 8 MB and 45 seconds. A page kept in a library
           keeps up to 100,000 words; a source for a cast, 12,000."""),
        t("""**Named.** Requests say who is asking: `AILawLab-agent/1.0` for agents and
           `AILawLab-crawler/1.0` for crawls, each with the lab's address and a contact, so that a
           site owner can see who is reading and how to reach the lab."""),
        t("""**Polite crawling.** A crawl stays on one website and goes one level deep. It fetches one
           page at a time at least two seconds apart, longer if robots.txt asks (up to 30 seconds),
           takes at most 100 pages, stops on HTTP 429 or 503 or after three failures in a row, and
           runs one at a time per site. It stores the licence or copyright notice it finds with
           every document."""),
        t("""**Recorded.** Searches with their results, documents read with their library versions,
           pages refused with their reasons, and every crawl with the rules it ran under and the
           outcome for each link are kept in the lab's database."""),
    ]),

    ("h1", "8  The lab's own interface"),
    ("p", t("""
The browser talks to the lab through JSON endpoints under `/ai_law_lab/api/`. They are the
portal's plumbing, not a public API: every one needs a signed-in, approved account, and changes
need the researcher or administrator role. The endpoints that reach outside are those for
database search and ingest (`/api/sources/search`, `/api/sources/ingest`), web links
(`/api/sources/links` and the two refresh endpoints), crawls (`/api/crawls`,
`/api/crawls/preview`) and cast sources (`/api/source/read`).
""")),

    ("h1", "9  Limitations and open questions"),
    ("bullets", [
        t("""**The AI-agent list is maintained by hand.** New agents appear regularly; the list in
           Table 3 should be reviewed each term."""),
        t("""**Machine-readable AI preferences are still being standardized.** The IETF AI
           Preferences working group [4] is defining a common vocabulary. The lab follows
           robots.txt, "noai" and TDMRep, which sites use today, and should adopt the IETF
           vocabulary once it settles."""),
        t("""**Permission to read is not a licence to reuse.** The licence notice a crawl records,
           and the address and retrieval time kept with every document, let a researcher check
           before relying on, quoting or publishing what was collected."""),
        t("""**Pages a person adds by hand are not checked** (Section 5). If the lab's practice
           should be stricter, the same check can be applied there."""),
    ]),

    ("h1", "Appendix: where it lives"),
    ("table", {
        "caption": t("""Table 4. Source files, under `src/ailawlab/`. Settings quoted in this note are
            defaults from `config.py`, each overridable by an `AILAWLAB_` environment variable.
            Tests: `tests/test_site_rules.py`, `tests/test_network_tools.py`,
            `tests/test_crawler.py`."""),
        "header": ["File", "What it does"],
        "mono_first": True,
        "rows": [
            ["sources.py", "The five legal databases: search and fetch"],
            ["network_tools.py, tools.py", "Agents' searching and reading; pacing; the fetch library"],
            ["research.py, experiments.py", "Run permissions; pages named in the brief"],
            ["site_rules.py", "robots.txt, AI agents, TDMRep and noai checks"],
            ["source_material.py", "Fetching, the public-address guard, size and time limits"],
            ["web_links.py", "Web pages kept in a library, and checking them again"],
            ["crawler.py", "Crawling a page's links into a library"],
            ["web/views.py, web/templates/run.html", "The run page: read online, and refused"],
        ],
        "widths": [0.42, 0.58],
    }),

    ("h1", "References"),
    ("refs", [
        "M. Koster, G. Illyes, H. Zeller and L. Sassman. *Robots Exclusion Protocol*. RFC 9309, "
        "Internet Engineering Task Force, September 2022. https://www.rfc-editor.org/rfc/rfc9309",
        "W3C Text and Data Mining Reservation Protocol Community Group. *TDM Reservation Protocol "
        "(TDMRep)*. Final Community Group Report. https://www.w3.org/community/tdmrep/",
        "Directive (EU) 2019/790 of the European Parliament and of the Council of 17 April 2019 on "
        "copyright and related rights in the Digital Single Market, Article 4.",
        "IETF AI Preferences (aipref) Working Group. https://datatracker.ietf.org/wg/aipref/about/",
        "Free Law Project. *CourtListener REST API, v4*. https://www.courtlistener.com/help/api/rest/",
        "Office of the Federal Register. *Federal Register API, v1*. "
        "https://www.federalregister.gov/developers/documentation/api/v1",
        "Office of the Federal Register. *eCFR API*. https://www.ecfr.gov/developers/documentation/api/v1",
        "US Government Publishing Office. *govinfo API*. https://api.govinfo.gov/docs/",
        "US Securities and Exchange Commission. *Accessing EDGAR Data*. "
        "https://www.sec.gov/os/accessing-edgar-data",
        "Brave Software. *Brave Search API*. https://api.search.brave.com/app/documentation",
        "17 U.S.C. § 105, Subject matter of copyright: United States Government works.",
    ]),
]


# ---------------------------------------------------------------- PDF


def _aligned(lines: list[str]) -> list[str]:
    """Procedure lines with their "->" outcomes in one column, and their "(...)" notes in another."""
    out = []
    for line in lines:
        if "->" in line:
            left, right = line.split("->", 1)
            line = f"{left.rstrip():<49} -> {right.strip()}"
        elif line.rstrip().endswith(")") and "  (" in line:
            left, note = line.rsplit("  (", 1)
            line = f"{left.rstrip():<52} ({note}"
        out.append(line)
    return out


def _pdf_markup(text: str) -> str:
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    # Code first, set aside, so a "*" inside it is not read as emphasis.
    code: list[str] = []

    def keep(m: re.Match) -> str:
        code.append(f'<font name="Mono">{m.group(1)}</font>')
        return f"\x00{len(code) - 1}\x00"

    text = re.sub(r"`(.+?)`", keep, text)
    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"\*(.+?)\*", r"<i>\1</i>", text)
    text = re.sub(r"\x00(\d+)\x00", lambda m: code[int(m.group(1))], text)
    return re.sub(r"(https?://[^\s<]+)", r'<link href="\1" color="#185fa5">\1</link>', text)


def build_pdf(path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        KeepTogether,
        Paragraph,
        Preformatted,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    fonts = Path("/usr/share/fonts/truetype/dejavu")
    faces = {"Serif": "DejaVuSerif.ttf", "Serif-Bold": "DejaVuSerif-Bold.ttf",
             "Serif-Italic": "DejaVuSerif-Italic.ttf", "Serif-BoldItalic": "DejaVuSerif-BoldItalic.ttf",
             "Mono": "DejaVuSansMono.ttf"}
    for name, file in faces.items():
        target = fonts / file
        if not target.exists():                      # italics ship in fonts-dejavu-extra
            target = fonts / ("DejaVuSerif-Bold.ttf" if "Bold" in name else "DejaVuSerif.ttf")
        pdfmetrics.registerFont(TTFont(name, str(target)))
    pdfmetrics.registerFontFamily("Serif", normal="Serif", bold="Serif-Bold",
                                  italic="Serif-Italic", boldItalic="Serif-BoldItalic")

    ink, muted = colors.HexColor("#1f1e1b"), colors.HexColor("#5f5b54")
    body = ParagraphStyle("body", fontName="Serif", fontSize=9.6, leading=13.4, alignment=TA_JUSTIFY,
                          textColor=ink, spaceAfter=6)
    styles = {
        "title": ParagraphStyle("title", parent=body, fontName="Serif-Bold", fontSize=16.5,
                                leading=20, alignment=TA_CENTER, spaceAfter=4),
        "subtitle": ParagraphStyle("subtitle", parent=body, fontName="Serif-Italic", fontSize=11,
                                   leading=14, alignment=TA_CENTER, spaceAfter=10),
        "author": ParagraphStyle("author", parent=body, fontSize=10, alignment=TA_CENTER, spaceAfter=1),
        "meta": ParagraphStyle("meta", parent=body, fontSize=8.6, alignment=TA_CENTER,
                               textColor=muted, spaceAfter=14),
        "abstract_head": ParagraphStyle("abstract_head", parent=body, fontName="Serif-Bold",
                                        fontSize=9.4, alignment=TA_LEFT, spaceAfter=2),
        "abstract": ParagraphStyle("abstract", parent=body, fontSize=9, leading=12.4),
        "h1": ParagraphStyle("h1", parent=body, fontName="Serif-Bold", fontSize=11.2, leading=14,
                             alignment=TA_LEFT, spaceBefore=10, spaceAfter=4, keepWithNext=True),
        "h2": ParagraphStyle("h2", parent=body, fontName="Serif-BoldItalic", fontSize=9.8,
                             leading=13, alignment=TA_LEFT, spaceBefore=6, spaceAfter=3,
                             keepWithNext=True),
        "caption": ParagraphStyle("caption", parent=body, fontSize=8.4, leading=11, alignment=TA_LEFT,
                                  textColor=muted, spaceBefore=2, spaceAfter=10),
        "cell": ParagraphStyle("cell", parent=body, fontSize=8.3, leading=10.6, alignment=TA_LEFT,
                               spaceAfter=0),
        "cell_mono": ParagraphStyle("cell_mono", parent=body, fontName="Mono", fontSize=7.9,
                                    leading=10.6, alignment=TA_LEFT, spaceAfter=0),
        "cell_head": ParagraphStyle("cell_head", parent=body, fontName="Serif-Bold", fontSize=8.3,
                                    leading=10.6, alignment=TA_LEFT, spaceAfter=0),
        "mono": ParagraphStyle("mono", fontName="Mono", fontSize=7.9, leading=10.4, textColor=ink),
        "bullet": ParagraphStyle("bullet", parent=body, leftIndent=12, bulletIndent=2, spaceAfter=4),
        "ref": ParagraphStyle("ref", parent=body, fontSize=8.4, leading=11.2, alignment=TA_LEFT,
                              leftIndent=20, firstLineIndent=-20, spaceAfter=3),
    }

    frame_width = letter[0] - 2 * inch
    story: list = [
        Paragraph(_pdf_markup(TITLE), styles["title"]),
        Paragraph(_pdf_markup(SUBTITLE), styles["subtitle"]),
        Paragraph(AUTHOR, styles["author"]),
        Paragraph(AFFILIATION, styles["author"]),
        Paragraph(DATELINE, styles["meta"]),
    ]
    abstract = Table([[Paragraph("Abstract", styles["abstract_head"])],
                      [Paragraph(_pdf_markup(ABSTRACT), styles["abstract"])]],
                     colWidths=[frame_width - 0.5 * inch])
    abstract.setStyle(TableStyle([
        ("LINEABOVE", (0, 0), (-1, 0), 0.6, muted), ("LINEBELOW", (0, -1), (-1, -1), 0.6, muted),
        ("LEFTPADDING", (0, 0), (-1, -1), 2), ("RIGHTPADDING", (0, 0), (-1, -1), 2),
        ("TOPPADDING", (0, 0), (-1, 0), 6), ("BOTTOMPADDING", (0, -1), (-1, -1), 4)]))
    story += [abstract, Spacer(1, 8)]

    grid = [("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#b9b4aa")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#efece6")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]

    for kind, value in CONTENT:
        if kind in ("h1", "h2"):
            story.append(Paragraph(_pdf_markup(value), styles[kind]))
        elif kind == "p":
            story.append(Paragraph(_pdf_markup(value), body))
        elif kind == "bullets":
            story += [Paragraph(_pdf_markup(b), styles["bullet"], bulletText="•") for b in value]
        elif kind == "refs":
            story += [Paragraph(f"[{i}]  {_pdf_markup(r)}", styles["ref"])
                      for i, r in enumerate(value, 1)]
        elif kind == "table":
            rows = [[Paragraph(_pdf_markup(h), styles["cell_head"]) for h in value["header"]]]
            for row in value["rows"]:
                rows.append([Paragraph(_pdf_markup(c), styles["cell_mono" if j == 0 and value.get("mono_first")
                                                         else "cell"]) for j, c in enumerate(row)])
            table = Table(rows, colWidths=[w * frame_width for w in value["widths"]], repeatRows=1)
            table.setStyle(TableStyle(grid))
            story.append(KeepTogether([table, Paragraph(_pdf_markup(value["caption"]), styles["caption"])]))
        elif kind == "algo":
            box = Table([[Preformatted("\n".join(_aligned(value["lines"])), styles["mono"])]],
                        colWidths=[frame_width])
            box.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#8f8a80")),
                                     ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f7f5f1")),
                                     ("LEFTPADDING", (0, 0), (-1, -1), 8),
                                     ("TOPPADDING", (0, 0), (-1, -1), 6),
                                     ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
            story.append(KeepTogether([box, Paragraph(_pdf_markup(value["caption"]), styles["caption"])]))

    def footer(canvas, doc) -> None:
        canvas.saveState()
        canvas.setFont("Serif", 7.8)
        canvas.setFillColor(muted)
        canvas.drawString(inch, 0.6 * inch, FOOTER)
        canvas.drawRightString(letter[0] - inch, 0.6 * inch, str(doc.page))
        canvas.restoreState()

    doc = SimpleDocTemplate(str(path), pagesize=letter, leftMargin=inch, rightMargin=inch,
                            topMargin=0.9 * inch, bottomMargin=0.9 * inch, title=TITLE,
                            author=f"{AUTHOR}, {AFFILIATION}", subject=SUBTITLE,
                            keywords="robots.txt, TDMRep, noai, legal databases, AI agents, web access")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


# ---------------------------------------------------------------- Word


def _docx_runs(paragraph, text: str, size: float | None = None, italic: bool = False) -> None:
    from docx.shared import Pt

    for part in re.split(r"(\*\*.+?\*\*|`.+?`|\*.+?\*)", text):
        if not part:
            continue
        run = paragraph.add_run()
        if part.startswith("**"):
            run.text, run.bold = part[2:-2], True
        elif part.startswith("`"):
            run.text = part[1:-1]
            run.font.name = "Consolas"
        elif part.startswith("*"):
            run.text, run.italic = part[1:-1], True
        else:
            run.text = part
        if italic:
            run.italic = True
        if size:
            run.font.size = Pt(size)


def build_docx(path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(1)
        section.top_margin = section.bottom_margin = Inches(0.9)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    for level, size in ((1, 12.5), (2, 11)):
        style = doc.styles[f"Heading {level}"]
        style.font.name, style.font.size = "Times New Roman", Pt(size)
        style.font.color.rgb = RGBColor(0x1F, 0x1E, 0x1B)
        style.font.italic = level == 2

    def centered(text: str, size: float, bold: bool = False, italic: bool = False, after: float = 2):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(after)
        run = p.add_run(text)
        run.bold, run.italic, run.font.size = bold, italic, Pt(size)
        return p

    centered(TITLE, 17, bold=True)
    centered(SUBTITLE, 11.5, italic=True, after=10)
    centered(AUTHOR, 10.5)
    centered(AFFILIATION, 10.5)
    centered(DATELINE, 9, after=14)

    head = doc.add_paragraph()
    head.add_run("Abstract").bold = True
    abstract = doc.add_paragraph()
    abstract.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _docx_runs(abstract, ABSTRACT, size=10)

    def table_cell_text(cell, text: str, bold: bool = False, mono: bool = False) -> None:
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        _docx_runs(p, f"**{text}**" if bold else text, size=9)
        if mono:
            for run in p.runs:
                run.font.name = "Consolas"

    def caption(text: str) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(3), Pt(10)
        _docx_runs(p, text, size=9, italic=True)

    for kind, value in CONTENT:
        if kind == "h1":
            doc.add_heading(value, level=1)
        elif kind == "h2":
            doc.add_heading(value, level=2)
        elif kind == "p":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _docx_runs(p, value)
        elif kind == "bullets":
            for item in value:
                _docx_runs(doc.add_paragraph(style="List Bullet"), item)
        elif kind == "refs":
            for i, ref in enumerate(value, 1):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.35)
                p.paragraph_format.first_line_indent = Inches(-0.35)
                p.paragraph_format.space_after = Pt(3)
                _docx_runs(p, f"[{i}]\t{ref}", size=9.5)
        elif kind == "table":
            table = doc.add_table(rows=1, cols=len(value["header"]))
            table.style = "Table Grid"
            for cell, text in zip(table.rows[0].cells, value["header"], strict=True):
                table_cell_text(cell, text, bold=True)
            for row in value["rows"]:
                cells = table.add_row().cells
                for j, (cell, text) in enumerate(zip(cells, row, strict=True)):
                    table_cell_text(cell, text, mono=j == 0 and value.get("mono_first", False))
            for row in table.rows:
                for cell, width in zip(row.cells, value["widths"], strict=True):
                    cell.width = Inches(6.5 * width)
            caption(value["caption"])
        elif kind == "algo":
            table = doc.add_table(rows=1, cols=1)
            table.style = "Table Grid"
            cell = table.rows[0].cells[0]
            cell.text = ""
            for k, line in enumerate(value["lines"]):
                p = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(_aligned([line])[0] or " ")
                run.font.name, run.font.size = "Consolas", Pt(8.5)
            caption(value["caption"])

    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = f"{FOOTER}    "
    for run in footer.runs:
        run.font.size = Pt(8)
    page = footer.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), tag)
        else:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        page._r.append(element)
    page.font.size = Pt(8)

    props = doc.core_properties
    props.title, props.subject = TITLE, SUBTITLE
    props.author, props.keywords = f"{AUTHOR}, {AFFILIATION}", "robots.txt, TDMRep, AI agents, data access"
    doc.save(str(path))


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    build_pdf(OUT / f"{STEM}.pdf")
    build_docx(OUT / f"{STEM}.docx")
    for suffix in ("pdf", "docx"):
        target = OUT / f"{STEM}.{suffix}"
        print(f"wrote {target.relative_to(ROOT)} ({target.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
