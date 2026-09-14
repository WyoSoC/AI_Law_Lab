"""Build the role-play moderation method note as PDF and Word from one content source.

The note documents graphs/roleplay.py and graphs/roleplay_policy.py. Keep the numbers here in
step with those files: every threshold quoted below is a constant there.

The generators are not dependencies of the app, so run this from a throwaway environment:

    uv venv /tmp/docenv && uv pip install --python /tmp/docenv/bin/python python-docx reportlab
    /tmp/docenv/bin/python docs/roleplay-moderation/build.py

Output goes to src/ailawlab/web/static/docs/, where the web portal links to it.
"""
from __future__ import annotations

import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "src" / "ailawlab" / "web" / "static" / "docs"
STEM = "roleplay-moderation"

TITLE = "Moderating Multi-Agent Legal Role-Play with a Small Local Language Model"
SUBTITLE = "Design and rationale of the AI Law Lab moderator"
AUTHOR = "Jian Gong"
AFFILIATION = "AI Law Lab, University of Wyoming"
DATELINE = "Method note, September 2026. Describes the software at commit aceecf2."
FOOTER = "AI Law Lab · Method note · Moderating multi-agent legal role-play"


def t(text: str) -> str:
    """Collapse a triple-quoted paragraph to one line."""
    return " ".join(text.split())


ABSTRACT = t("""
Role-play between language-model agents is a promising instrument for legal education and
research, but multi-agent exchanges fail in characteristic ways: agents repeat themselves,
drift out of character, converge too early, or let one voice dominate. We describe the
moderation logic of the AI Law Lab's role-play mode, which runs an 8-billion-parameter model on
local hardware. The design separates judgement from rules: the model answers four structured
questions about each turn, while code decides who speaks, when an impasse warrants
intervention, and when a scene may end. Participants keep private negotiation ledgers and
bottom lines that neither the moderator nor the other participants see. Each rule is grounded
in prior work on LLM negotiation, multi-party turn-taking, multi-agent debate and mediation
practice. A live 24-turn, three-party mediation illustrates the mechanisms and exposes two
limitations: routing direct questions can override turn balance, and interventions can be
addressed to a neutral participant.
""")

# Blocks: ("h1", text) ("p", text) ("table", spec) ("algo", spec) ("bullets", [text]) ("refs", [text])
# Inline markup: **bold**, *italic*, `code`. URLs are linked in the PDF.
CONTENT: list[tuple] = [
    ("h1", "1  Introduction"),
    ("p", t("""
Legal education has long used simulated negotiations, mediations and hearings. Language-model
agents make such simulations cheap to run at scale and, when every model call is logged,
auditable, which makes them attractive both for teaching and for research on how models reason
about legal disputes. The AI Law Lab runs these simulations on university hardware with an
8-billion-parameter model, so that unpublished or sensitive material never leaves campus.
""")),
    ("p", t("""
Exchanges between language-model agents fail in recognizable ways. Negotiating agents without
a plan re-propose the same deal round after round [1]; simulations can look varied yet end in
impasse almost every time, far from human base rates [2]; without turn-taking structure,
individual agents monopolize the floor or pose question after question [3]; agents abandon
positions under peer pressure and collapse into premature agreement [4], [5]; and
instruction-following drifts measurably within about eight rounds of dialog [6]. A small local
model is, if anything, more exposed to these failures than the larger models studied in that
work.
""")),
    ("p", t("""
This note describes how the lab's role-play moderator addresses them. Its central design
choice is to divide the work: the model supplies judgement about each turn in a constrained,
structured form, and deterministic code applies the rules that decide who speaks, when to
intervene and when a scene may end. Section 3 outlines the system and what each party may see,
Section 4 specifies the moderation rules, Section 5 the participant-side mechanisms they rely
on, and Section 6 the assessment. Section 7 walks through a live run and Section 8 sets out
limitations.
""")),

    ("h1", "2  Background"),
    ("p", t("""
Conversation analysis describes how speakers allocate turns: the current speaker may select
the next, most visibly by producing the first part of an adjacency pair such as a question,
and otherwise participants select themselves [7]. Nonomura and Mori applied this ordering to
multi-party dialog between LLM agents and found that letting the current speaker select the
next, before any self-selection, reduced breakdowns such as one agent monopolizing the floor or
asking question after question [3].
""")),
    ("p", t("""
Principled negotiation separates positions from the interests behind them and asks each party
to know its best alternative to a negotiated agreement (BATNA), which tells it when to walk
away [8]. Mediators manage stalled talks with process tools, among them reality testing each
side's assessment of its case, reframing and refocusing the discussion, and caucuses [9]. The
moderator borrows these tools; the participant specification borrows the preparation.
""")),

    ("h1", "3  System overview"),
    ("p", t("""
A role-play is a LangGraph state machine over a shared transcript. In each cycle the moderator
judges the latest turn and chooses the next speaker; the chosen participant first updates a
private ledger and then speaks; the cycle repeats until the scene ends, after which an assessor
evaluates the whole exchange. All model calls go to gemma4 (8 billion parameters, 4-bit
quantized) served by Ollama on two NVIDIA DGX Spark units. Calls whose output code must read are
constrained by a JSON schema. Every call, decision and intervention is written to an
append-only trace.
""")),
    ("p", t("""
Participants are specified by a role, an objective, and optional background, demeanor,
behavioral tendencies and priorities, plus two private fields: a bottom line and confidential
facts. Table 1 shows what each party can see. The separation protects validity: a moderator
that saw bottom lines could steer the parties toward their zone of agreement, and a participant
that saw another's ledger would no longer be negotiating.
""")),
    ("table", {
        "caption": "Table 1. What each party sees during a run.",
        "header": ["Information", "Speaking participant", "Other participants", "Moderator",
                   "Assessor"],
        "rows": [
            ["Scenario", "yes", "yes", "yes", "yes"],
            ["Public transcript, including moderator statements", "yes", "yes",
             "latest 8 entries", "yes"],
            ["Participants' names and roles", "all", "all", "all", "names only"],
            ["Objective", "own", "own", "all", "no"],
            ["Background, demeanor, tendencies, priorities", "own", "own", "no", "no"],
            ["Bottom line and confidential facts", "own", "own", "no", "all"],
            ["Private negotiation ledger", "own", "own", "no", "no"],
            ["Moderator's private note (e.g. answer a question)", "if addressed", "no",
             "writes it", "no"],
        ],
        "widths": [0.34, 0.17, 0.17, 0.17, 0.15],
    }),

    ("h1", "4  Moderation logic"),
    ("h2", "4.1  Structured judgement"),
    ("p", t("""
Before every turn after the first, the moderator is shown the scenario, each participant's
role and objective, how often each has spoken, the seven preceding transcript entries shortened
to 120 words each, and the latest entry in full. It must answer four questions in a fixed
schema (Table 2). Reasoning is disabled, temperature is 0 and output is capped at 300 tokens:
the task is classification, not analysis. Because `next_speaker` and `question_for` are
constrained to the participants' identifiers, no free-text parsing is needed. The opening turn
goes to the first listed participant without a model call.
""")),
    ("table", {
        "caption": "Table 2. The moderator's structured judgement. A fifth field, reason, is kept in the trace.",
        "header": ["Field", "Question put to the model", "Used by"],
        "rows": [
            ["question_for", ("Did the latest speaker put a direct question to, or demand an answer "
             "from, one specific other participant? (The person asked, never the asker, or "
             "\"nobody\".)"), "Rule 1, Section 4.2"],
            ["stalled", ("Over the last several turns, are the parties restating the same positions "
             "without any new proposal, concession or question?"), "Interventions, Section 4.3"],
            ["concluded", ("Have the parties clearly agreed on terms, or has someone clearly and "
             "finally walked away?"), "Ending, Section 4.4"],
            ["next_speaker", ("Whose turn should come next for the exchange to be productive and "
             "fair?"), "Rules 2 to 4, Section 4.2"],
        ],
        "widths": [0.17, 0.60, 0.23],
        "mono_first": True,
    }),
    ("h2", "4.2  Speaker selection"),
    ("p", t("""
Code turns the judgement into a choice with Algorithm 1. The previous speaker is never
eligible. A direct question is answered first, following the adjacency-pair ordering of [7]
that reduced breakdowns in [3]. The moderator's own suggestion stands unless it would put that
participant two or more turns ahead of the quietest eligible participant, which guards against
the monopolization that free self-selection produced in [3]. Ties among the quietest go to the
next seat after the previous speaker. A participant chosen to answer a question also receives a
private note to answer it before anything else.
""")),
    ("algo", {
        "caption": "Algorithm 1. Choosing the next speaker (pick_speaker in roleplay_policy.py).",
        "lines": [
            "input   participants in seating order; previous speaker p;",
            "        turns spoken c(i); judgement: question_for q, next_speaker s",
            "",
            "E  <-  participants other than p",
            "m  <-  participant in E with the fewest turns (ties: next seat after p)",
            "",
            "1  if q is in E                       return q   answer the question",
            "2  if s is not in E                   return m   unusable suggestion",
            "3  if c(s) - min{ c(i) : i in E } >= 2",
            "                                      return m   balance the floor",
            "4  return s                                      moderator's choice",
        ],
    }),
    ("h2", "4.3  Impasse interventions"),
    ("p", t("""
When the judgement reports a stall, the moderator intervenes provided at least max(4, 2n) turns
have passed since its previous intervention, where n is the number of participants. The
cool-down gives the parties room to respond before it steps in again. Successive interventions
rotate through three mediation techniques (Table 3). Each is addressed by name to the
participant who speaks next, entered in the public transcript as a statement made in joint
session, and repeated to that participant as a private note.
""")),
    ("table", {
        "caption": "Table 3. Intervention techniques, used in rotation. Statements abridged.",
        "header": ["Technique", "Statement to the next speaker", "Basis"],
        "rows": [
            ["Reframe", ("\"Before restating anything, say what you think the other side needs "
             "underneath its demand, and propose one option that could serve both sides.\""),
             "Interests over positions [8]; reframing [9]"],
            ["Narrow", ("\"Pick the single open issue closest to agreement and make a concrete, "
             "specific proposal on that issue alone.\""), "Refocusing the discussion [9]"],
            ["Reality test", ("\"Consider plainly what happens to your side if there is no agreement "
             "at all. Then either make a conditional offer or say clearly that you are prepared "
             "to walk away.\""), "BATNA [8]; reality testing [9]"],
        ],
        "widths": [0.15, 0.58, 0.27],
    }),
    ("h2", "4.4  Ending a scene"),
    ("p", t("""
A scene ends at the configured maximum number of turns, or earlier when the judgement reports
it concluded and every participant has spoken at least twice. The floor exists because
multi-agent exchanges drift toward premature consensus and abandon positions under social
pressure [4], [5]. A conclusion judged too early is written to the trace and the exchange
continues.
""")),

    ("h1", "5  Participant-side mechanisms"),
    ("p", t("""
Before speaking, the selected participant updates a private negotiation ledger with seven
fields: concessions made, concessions received, where things stand, its read on the others,
open issues, arguments already made, and its next move. The update is a schema-constrained call
with reasoning disabled that carries the previous ledger forward; a malformed update keeps the
previous ledger. The design follows two findings: long-term planning notes kept negotiating
agents from re-proposing similar deals [1], and in simulated negotiations an agent-authored
ledger of concessions and open issues was the one intervention that moved outcomes off
near-universal impasse, where extra output tokens, free-form reflection and reasoning modes did
not [2]. The arguments-already-made field targets the repetition seen in the lab's earlier runs.
""")),
    ("p", t("""
The reply prompt then presents, in order, what the participant has not yet responded to, its
ledger, any note from the moderator, and a short character reminder restating its identity,
objective, private bottom line and demeanor. The reminder comes last because adherence to
initial instructions decays over long dialogs [6]. Replies use the model's reasoning mode with
an output budget proportional to the word limit, and a reply exhausted by reasoning before
producing any text is retried without it. Each participant's memory is scoped to that
participant and stores every statement once, as what it heard since its last turn and what it
said, compacting the oldest material into summaries beyond a token budget scaled to the turn
length.
""")),

    ("h1", "6  Assessment"),
    ("p", t("""
After the scene an assessor receives the transcript together with the participants' bottom
lines and confidential facts, which only it is shown, and answers five questions citing turn
numbers: the outcome and its terms; which positions moved and why, and where the exchange
repeated itself; whether anyone conceded past a bottom line or disclosed confidential
information; any legal proposition that appears unsupported; and whether concessions and
resistance were plausible for real counsel. A transcript above 60,000 tokens is split into
segments of at most 24,000 tokens that are summarized concurrently, then assessed with the final
six entries verbatim; a 100-turn scene at 1,000 words a turn would otherwise exceed the model's
128,000-token context.
""")),

    ("h1", "7  Illustrative run"),
    ("p", t("""
Table 4 summarizes a live run from September 2026. Its scenario, drafted by the lab from a news
report of Wyoming Supreme Court arguments over state trust-land gravel leases, set up a
mediation between a mediator, an Assistant Attorney General representing the State Board of Land
Commissioners, and the manager of a gravel operator. The names of real people and private
organizations in the source were replaced before drafting. The run allowed up to 100 turns of up
to 1,000 words each.
""")),
    ("table", {
        "caption": "Table 4. The run in numbers.",
        "header": ["Measure", "Value"],
        "rows": [
            ["Length", "24 turns in 12 min 2 s"],
            ["How it ended", ("The moderator judged the scene concluded after the operator's "
             "representative withdrew; the assessment found no agreement")],
            ["Speaker selection", ("16 answering a direct question, 5 moderator's choice, 2 balancing "
             "speaking time; the opening turn by rule")],
            ["Interventions", "Reframe after turn 6, narrow after turn 15, reality test after turn 21"],
            ["Premature conclusions held back", "1, after turn 13"],
            ["Turns by participant", "State's counsel 11, operator 7, mediator 6"],
            ["Reply length", "mean 454 words, maximum 593"],
            ["Mean model time per call", ("reply 17.0 s (949 output tokens), ledger 3.7 s, moderator "
             "2.4 s, assessment 35.1 s")],
            ["Unparseable decisions or ledgers; truncated replies", "0; 0"],
        ],
        "widths": [0.32, 0.68],
    }),
    ("p", t("""
The mechanisms behaved as specified. Direct questions were routed to the participants asked, all
three techniques fired in rotation with the cool-down respected, a conclusion judged after turn
13 was held back because not every participant had yet spoken twice, and the scene ended when
one party walked away. A shorter check with three participants, a 150-word limit and a nine-turn
cap behaved likewise, with three question-routed turns and one reframe intervention. These are
single illustrative runs, not an evaluation.
""")),

    ("h1", "8  Limitations"),
    ("bullets", [
        t("""**Question routing can override balance.** Rule 1 precedes the balance rule, and most
        questions in the run were put to the State's counsel, who took 11 of 24 turns. Capping
        consecutive question-routed turns for one participant, or applying the balance rule to
        questions beyond a wider margin, would address this."""),
        t("""**Interventions can reach the wrong participant.** They are addressed to whoever speaks
        next, which twice in the run was the neutral mediator, including a reality test that
        asked the mediator to weigh “your side”. Interventions should be addressed to a
        party, or phrased for a neutral."""),
        t("""**The judgements are unvalidated.** Question, stall and conclusion detection come from an
        8-billion-parameter model at temperature 0 and have not been compared with human
        coders."""),
        t("""**The parameters are heuristic.** The balance margin, the cool-down and the fixed rotation
        have not been ablated."""),
        t("""**Outcome fidelity is untested.** Varied-looking simulations can still produce
        unrealistic outcome distributions [2]; this configuration has not been checked against
        real negotiation base rates."""),
        t("""**Scope.** One model on one hardware configuration; the timings do not transfer."""),
    ]),

    ("h1", "9  Conclusion"),
    ("p", t("""
Dividing moderation into structured model judgement and code-enforced rules lets a small local
model run legal role-plays in which every decision can be inspected afterwards. The live run
shows the rules working as designed and points to the next revisions: bounding question routing,
and addressing interventions only to parties.
""")),

    ("h1", "Acknowledgments"),
    ("p", t("""
The moderation logic was implemented, and this note drafted, with the assistance of Claude
(Anthropic). The participant names in the run of Section 7 are invented.
""")),

    ("h1", "References"),
    ("refs", [
        t("""S. Abdelnabi, A. Gomaa, S. Sivaprasad, L. Schönherr, and M. Fritz, “Cooperation,
        Competition, and Maliciousness: LLM-Stakeholders Interactive Negotiation,” arXiv
        preprint arXiv:2309.17234, 2023. https://arxiv.org/abs/2309.17234"""),
        t("""S. Andric, “Diversity Without Fidelity: A Solver-Sampler Mismatch in Multi-Agent LLM
        Negotiation Simulation,” arXiv preprint arXiv:2604.11840, 2026.
        https://arxiv.org/abs/2604.11840"""),
        t("""R. Nonomura and H. Mori, “Who Speaks Next? Multi-party AI Discussion Leveraging the
        Systematics of Turn-taking in Murder Mystery Games,” *Frontiers in Artificial
        Intelligence*, 2025, doi: 10.3389/frai.2025.1582287. Preprint arXiv:2412.04937.
        https://arxiv.org/abs/2412.04937"""),
        t("""B. Yao, C. Shang, W. Du, J. He, R. Lian, Y. Zhang, H. Su, S. Swamy, and Y. Qi,
        “Peacemaker or Troublemaker: How Sycophancy Shapes Multi-Agent Debate,” arXiv preprint
        arXiv:2509.23055, 2025. https://arxiv.org/abs/2509.23055"""),
        t("""A. Wynn, H. Satija, and G. Hadfield, “Talk Isn't Always Cheap: Understanding Failure
        Modes in Multi-Agent Debate,” ICML 2025 Workshop on Multi-Agent Systems, 2025. Preprint
        arXiv:2509.05396. https://arxiv.org/abs/2509.05396"""),
        t("""K. Li, T. Liu, N. Bashkansky, D. Bau, F. Viégas, H. Pfister, and M. Wattenberg,
        “Measuring and Controlling Instruction (In)Stability in Language Model Dialogs,” in
        *Proc. Conference on Language Modeling (COLM)*, 2024. Preprint arXiv:2402.10962.
        https://arxiv.org/abs/2402.10962"""),
        t("""H. Sacks, E. A. Schegloff, and G. Jefferson, “A Simplest Systematics for the
        Organization of Turn-Taking for Conversation,” *Language*, vol. 50, no. 4, pp. 696–735,
        1974."""),
        t("""R. Fisher, W. Ury, and B. Patton, *Getting to Yes: Negotiating Agreement Without Giving
        In*, 3rd ed. New York, NY, USA: Penguin, 2011."""),
        t("""J. S. Johnsen, “Managing the Mediation Process,” *AAA Mediation Magazine*, Mar. 23,
        2026. https://mediationmagazine.adr.org/managing-the-mediation-process/"""),
    ]),
]


# ---------------------------------------------------------------- PDF


def _pdf_markup(text: str) -> str:
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"\*(.+?)\*", r"<i>\1</i>", text)
    text = re.sub(r"`(.+?)`", r'<font name="Mono">\1</font>', text)
    return re.sub(r"(https?://[^\s<]+)", r'<link href="\1" color="#185fa5">\1</link>', text)


def build_pdf(path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        KeepTogether,
        Paragraph,
        Preformatted,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    fonts = Path("/usr/share/fonts/truetype/dejavu")
    faces = {"Serif": "DejaVuSerif.ttf", "Serif-Bold": "DejaVuSerif-Bold.ttf",
             "Serif-Italic": "DejaVuSerif-Italic.ttf", "Serif-BoldItalic": "DejaVuSerif-BoldItalic.ttf",
             "Mono": "DejaVuSansMono.ttf"}
    for name, file in faces.items():
        target = fonts / file
        if not target.exists():                      # italics ship in fonts-dejavu-extra
            target = fonts / ("DejaVuSerif-Bold.ttf" if "Bold" in name else "DejaVuSerif.ttf")
        pdfmetrics.registerFont(TTFont(name, str(target)))
    pdfmetrics.registerFontFamily("Serif", normal="Serif", bold="Serif-Bold",
                                  italic="Serif-Italic", boldItalic="Serif-BoldItalic")

    ink, muted = colors.HexColor("#1f1e1b"), colors.HexColor("#5f5b54")
    body = ParagraphStyle("body", fontName="Serif", fontSize=9.6, leading=13.4, alignment=TA_JUSTIFY,
                          textColor=ink, spaceAfter=6)
    styles = {
        "title": ParagraphStyle("title", parent=body, fontName="Serif-Bold", fontSize=16.5,
                                leading=20, alignment=TA_CENTER, spaceAfter=4),
        "subtitle": ParagraphStyle("subtitle", parent=body, fontName="Serif-Italic", fontSize=11,
                                   leading=14, alignment=TA_CENTER, spaceAfter=10),
        "author": ParagraphStyle("author", parent=body, fontSize=10, alignment=TA_CENTER, spaceAfter=1),
        "meta": ParagraphStyle("meta", parent=body, fontSize=8.6, alignment=TA_CENTER,
                               textColor=muted, spaceAfter=14),
        "abstract_head": ParagraphStyle("abstract_head", parent=body, fontName="Serif-Bold",
                                        fontSize=9.4, alignment=TA_LEFT, spaceAfter=2),
        "abstract": ParagraphStyle("abstract", parent=body, fontSize=9, leading=12.4),
        "h1": ParagraphStyle("h1", parent=body, fontName="Serif-Bold", fontSize=11.2, leading=14,
                             alignment=TA_LEFT, spaceBefore=10, spaceAfter=4, keepWithNext=True),
        "h2": ParagraphStyle("h2", parent=body, fontName="Serif-BoldItalic", fontSize=9.8,
                             leading=13, alignment=TA_LEFT, spaceBefore=6, spaceAfter=3,
                             keepWithNext=True),
        "caption": ParagraphStyle("caption", parent=body, fontSize=8.4, leading=11, alignment=TA_LEFT,
                                  textColor=muted, spaceBefore=2, spaceAfter=10),
        "cell": ParagraphStyle("cell", parent=body, fontSize=8.3, leading=10.6, alignment=TA_LEFT,
                               spaceAfter=0),
        "cell_mono": ParagraphStyle("cell_mono", parent=body, fontName="Mono", fontSize=7.9,
                                    leading=10.6, alignment=TA_LEFT, spaceAfter=0),
        "cell_head": ParagraphStyle("cell_head", parent=body, fontName="Serif-Bold", fontSize=8.3,
                                    leading=10.6, alignment=TA_LEFT, spaceAfter=0),
        "mono": ParagraphStyle("mono", fontName="Mono", fontSize=7.9, leading=10.4, textColor=ink),
        "bullet": ParagraphStyle("bullet", parent=body, leftIndent=12, bulletIndent=2, spaceAfter=4),
        "ref": ParagraphStyle("ref", parent=body, fontSize=8.4, leading=11.2, alignment=TA_LEFT,
                              leftIndent=20, firstLineIndent=-20, spaceAfter=3),
    }

    frame_width = letter[0] - 2 * inch
    story: list = [
        Paragraph(_pdf_markup(TITLE), styles["title"]),
        Paragraph(_pdf_markup(SUBTITLE), styles["subtitle"]),
        Paragraph(AUTHOR, styles["author"]),
        Paragraph(AFFILIATION, styles["author"]),
        Paragraph(DATELINE, styles["meta"]),
    ]
    abstract = Table([[Paragraph("Abstract", styles["abstract_head"])],
                      [Paragraph(_pdf_markup(ABSTRACT), styles["abstract"])]],
                     colWidths=[frame_width - 0.5 * inch])
    abstract.setStyle(TableStyle([
        ("LINEABOVE", (0, 0), (-1, 0), 0.6, muted), ("LINEBELOW", (0, -1), (-1, -1), 0.6, muted),
        ("LEFTPADDING", (0, 0), (-1, -1), 2), ("RIGHTPADDING", (0, 0), (-1, -1), 2),
        ("TOPPADDING", (0, 0), (-1, 0), 6), ("BOTTOMPADDING", (0, -1), (-1, -1), 4)]))
    story += [abstract, Spacer(1, 8)]

    grid = [("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#b9b4aa")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#efece6")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]

    for kind, value in CONTENT:
        if kind in ("h1", "h2"):
            story.append(Paragraph(_pdf_markup(value), styles[kind]))
        elif kind == "p":
            story.append(Paragraph(_pdf_markup(value), body))
        elif kind == "bullets":
            story += [Paragraph(_pdf_markup(b), styles["bullet"], bulletText="•") for b in value]
        elif kind == "refs":
            story += [Paragraph(f"[{i}]  {_pdf_markup(r)}", styles["ref"])
                      for i, r in enumerate(value, 1)]
        elif kind == "table":
            rows = [[Paragraph(_pdf_markup(h), styles["cell_head"]) for h in value["header"]]]
            for row in value["rows"]:
                rows.append([Paragraph(_pdf_markup(c), styles["cell_mono" if j == 0 and value.get("mono_first")
                                                         else "cell"]) for j, c in enumerate(row)])
            table = Table(rows, colWidths=[w * frame_width for w in value["widths"]], repeatRows=1)
            table.setStyle(TableStyle(grid))
            story.append(KeepTogether([table, Paragraph(_pdf_markup(value["caption"]), styles["caption"])]))
        elif kind == "algo":
            box = Table([[Preformatted("\n".join(value["lines"]), styles["mono"])]],
                        colWidths=[frame_width])
            box.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#8f8a80")),
                                     ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f7f5f1")),
                                     ("LEFTPADDING", (0, 0), (-1, -1), 8),
                                     ("TOPPADDING", (0, 0), (-1, -1), 6),
                                     ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
            story.append(KeepTogether([box, Paragraph(_pdf_markup(value["caption"]), styles["caption"])]))

    def footer(canvas, doc) -> None:
        canvas.saveState()
        canvas.setFont("Serif", 7.8)
        canvas.setFillColor(muted)
        canvas.drawString(inch, 0.6 * inch, FOOTER)
        canvas.drawRightString(letter[0] - inch, 0.6 * inch, str(doc.page))
        canvas.restoreState()

    doc = SimpleDocTemplate(str(path), pagesize=letter, leftMargin=inch, rightMargin=inch,
                            topMargin=0.9 * inch, bottomMargin=0.9 * inch, title=TITLE,
                            author=f"{AUTHOR}, {AFFILIATION}", subject=SUBTITLE,
                            keywords="multi-agent, role-play, moderation, legal education, LLM")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


# ---------------------------------------------------------------- Word


def _docx_runs(paragraph, text: str, size: float | None = None, italic: bool = False) -> None:
    from docx.shared import Pt

    for part in re.split(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)", text):
        if not part:
            continue
        run = paragraph.add_run()
        if part.startswith("**"):
            run.text, run.bold = part[2:-2], True
        elif part.startswith("`"):
            run.text = part[1:-1]
            run.font.name = "Consolas"
        elif part.startswith("*"):
            run.text, run.italic = part[1:-1], True
        else:
            run.text = part
        if italic:
            run.italic = True
        if size:
            run.font.size = Pt(size)


def build_docx(path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(1)
        section.top_margin = section.bottom_margin = Inches(0.9)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    for level, size in ((1, 12.5), (2, 11)):
        style = doc.styles[f"Heading {level}"]
        style.font.name, style.font.size = "Times New Roman", Pt(size)
        style.font.color.rgb = RGBColor(0x1F, 0x1E, 0x1B)
        style.font.italic = level == 2

    def centered(text: str, size: float, bold: bool = False, italic: bool = False, after: float = 2):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(after)
        run = p.add_run(text)
        run.bold, run.italic, run.font.size = bold, italic, Pt(size)
        return p

    centered(TITLE, 17, bold=True)
    centered(SUBTITLE, 11.5, italic=True, after=10)
    centered(AUTHOR, 10.5)
    centered(AFFILIATION, 10.5)
    centered(DATELINE, 9, after=14)

    head = doc.add_paragraph()
    head.add_run("Abstract").bold = True
    abstract = doc.add_paragraph()
    abstract.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _docx_runs(abstract, ABSTRACT, size=10)

    def table_cell_text(cell, text: str, bold: bool = False, mono: bool = False) -> None:
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        _docx_runs(p, f"**{text}**" if bold else text, size=9)
        if mono:
            for run in p.runs:
                run.font.name = "Consolas"

    def caption(text: str) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(3), Pt(10)
        _docx_runs(p, text, size=9, italic=True)

    for kind, value in CONTENT:
        if kind == "h1":
            doc.add_heading(value, level=1)
        elif kind == "h2":
            doc.add_heading(value, level=2)
        elif kind == "p":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _docx_runs(p, value)
        elif kind == "bullets":
            for item in value:
                _docx_runs(doc.add_paragraph(style="List Bullet"), item)
        elif kind == "refs":
            for i, ref in enumerate(value, 1):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.35)
                p.paragraph_format.first_line_indent = Inches(-0.35)
                p.paragraph_format.space_after = Pt(3)
                _docx_runs(p, f"[{i}]\t{ref}", size=9.5)
        elif kind == "table":
            table = doc.add_table(rows=1, cols=len(value["header"]))
            table.style = "Table Grid"
            for cell, text in zip(table.rows[0].cells, value["header"], strict=True):
                table_cell_text(cell, text, bold=True)
            for row in value["rows"]:
                cells = table.add_row().cells
                for j, (cell, text) in enumerate(zip(cells, row, strict=True)):
                    table_cell_text(cell, text, mono=j == 0 and value.get("mono_first", False))
            for row in table.rows:
                for cell, width in zip(row.cells, value["widths"], strict=True):
                    cell.width = Inches(6.5 * width)
            caption(value["caption"])
        elif kind == "algo":
            table = doc.add_table(rows=1, cols=1)
            table.style = "Table Grid"
            cell = table.rows[0].cells[0]
            cell.text = ""
            for k, line in enumerate(value["lines"]):
                p = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(line or " ")
                run.font.name, run.font.size = "Consolas", Pt(8.5)
            caption(value["caption"])

    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = f"{FOOTER}    "
    for run in footer.runs:
        run.font.size = Pt(8)
    page = footer.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), tag)
        else:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        page._r.append(element)
    page.font.size = Pt(8)

    props = doc.core_properties
    props.title, props.subject = TITLE, SUBTITLE
    props.author, props.keywords = f"{AUTHOR}, {AFFILIATION}", "multi-agent, role-play, moderation"
    doc.save(str(path))


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    build_pdf(OUT / f"{STEM}.pdf")
    build_docx(OUT / f"{STEM}.docx")
    for suffix in ("pdf", "docx"):
        target = OUT / f"{STEM}.{suffix}"
        print(f"wrote {target.relative_to(ROOT)} ({target.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
